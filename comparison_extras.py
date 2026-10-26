import io
import re

# ─────────────────────────────────────────────────────────────
# 비교표 추가 기능 — app.py에서 불러 씁니다. (API 키 없이 동작)
# 1) 용법·용량 간단/일반/상세 원문 발췌와 버전 추천 (규칙 기반)
#    기준에 맞춘 최종 요약은 Excel을 받아 Claude(구독)가 채웁니다 → fill_dosage.py
# 2) 사람이 채우는 빈 행
# 3) 조회 실패 / 자료 없음 표시
# 4) 비교표 Excel 파일 만들기 (용법·용량 원문 시트 포함)
# streamlit 없이도 불러올 수 있게 만들어 따로 시험할 수 있습니다.
# ─────────────────────────────────────────────────────────────

STATUS_FAILED = "조회 실패"
STATUS_EMPTY = "자료 없음"
STATUS_VALUES = {STATUS_FAILED, STATUS_EMPTY}

# 사람이 직접 채우는 행 (비교표 예시 모음.pptx 기준). 원내코드는 대외비라 항상 빈칸으로 만듭니다.
MANUAL_HEAD_COLUMNS = ["원내코드"]
MANUAL_TAIL_COLUMNS = ["code open", "처방건수 (입원, 월평균)", "처방건수 (외래, 월평균)", "비급여"]
MANUAL_COLUMNS = MANUAL_HEAD_COLUMNS + MANUAL_TAIL_COLUMNS
APPLICANT_MARK = "[신청의약품]"

DOSAGE_LEVELS = ["간단", "일반", "상세"]
DOSAGE_LEVEL_KEYS = {"간단": "simple", "일반": "general", "상세": "detailed"}
DOSAGE_CHOICES = ["원문", "자동(추천)", "간단", "일반", "상세"]

# Excel 안에서 용법·용량 원문을 담는 시트. fill_dosage.py가 이 시트를 읽습니다.
SOURCE_SHEET = "용법용량 원문"
SOURCE_HEADERS = ["비교표 열 이름", "품목코드", "허가제품명", "효능효과 원문", "용법용량 원문", "비교표 시트"]
SOURCE_KEYS = ["column", "id", "name", "indications_text", "dosage_text", "sheet"]
COMPARISON_SHEET = "비교표"  # 한 번에 하나 만들 때. 신청번호별로 만들 때는 '신청1', '신청2' … 시트


def _clean(text):
    if text is None:
        return ""
    return re.sub(r"\s+", " ", str(text)).strip()


def is_status(value):
    return _clean(value) in STATUS_VALUES


def mark_missing(value, failed):
    """빈 값을 '조회 실패'(API 호출이 실패함) 또는 '자료 없음'(호출은 됐지만 내용이 없음)으로 바꿉니다."""
    if _clean(value):
        return value
    return STATUS_FAILED if failed else STATUS_EMPTY


# ── 용법·용량 (규칙 기반) ───────────────────────────────────

def recommend_level_by_rule(texts):
    """원문 단어로 버전을 추천합니다. 비교표 안에서 가장 복잡한 약 기준으로 고릅니다."""
    detailed_pattern = r"(\d+\s*kg\s*(이상|미만|초과|이하)|[≥≤<>]\s*\d+\s*kg|소아|크레아티닌|CrCl|신장애|신기능|휴약|주기|유도|유지요법)"
    general_pattern = r"(증량|개시|초회|초기|유지용량|공복|식전|식후|식사|정맥|점적|분간|시간에 걸쳐|피하|근육|분할)"
    level, reason = "간단", "모든 제품이 고정 용량이라 투여 횟수와 1회 용량만 비교하면 됩니다."
    for text in texts:
        if re.search(detailed_pattern, text or ""):
            return "상세", "체중·연령별 용량, 유도/유지, 투여 주기 또는 신장애 용량 조절이 있는 제품이 있습니다."
        if re.search(general_pattern, text or ""):
            level, reason = "일반", "개시·증량 단계나 투여경로·복용 조건이 제품마다 다릅니다."
    return level, reason


def _split_segments(text):
    parts = re.split(r"(?<=[.。])\s+|\s(?=\d{1,2}\.\s)|\s(?=<[^<>]{1,20}>)", _clean(text))
    return [part.strip() for part in parts if part.strip()]


def rule_based_versions(text):
    """원문을 잘라 세 버전을 만듭니다. 새 내용은 만들지 않고 앞부분만 발췌합니다."""
    text = _clean(text)
    if not text or is_status(text):
        return {"simple": text, "general": text, "detailed": text}
    segments = _split_segments(text)

    def take(limit_chars, limit_segments):
        picked = []
        for segment in segments[:limit_segments]:
            if picked and len(" ".join(picked + [segment])) > limit_chars:
                break
            picked.append(segment)
        joined = "\n".join(picked) or text[:limit_chars]
        return joined if len(joined) <= limit_chars else joined[:limit_chars].rstrip() + "…"

    return {"simple": take(90, 1), "general": take(280, 3), "detailed": "\n".join(segments)}


def _numbers(text):
    return set(re.findall(r"\d+(?:\.\d+)?", (text or "").replace(",", "")))


def unsupported_numbers(summary, source):
    """요약에 있는데 원문에는 없는 숫자를 찾습니다. 요약이 숫자를 지어냈는지 약사가 확인할 수 있게 합니다."""
    return sorted(_numbers(summary) - _numbers(source), key=lambda value: float(value))


def build_dosage_summaries(products):
    """제품 목록의 용법·용량 발췌 세 버전과 추천 버전을 만듭니다.

    products: [{"id", "name", "dosage_text", "indications_text"}]
    반환: {"recommended_level", "reason", "source", "items": {id: {...}}, "warnings": [...]}
    """
    usable = [p for p in products if _clean(p.get("dosage_text")) and not is_status(p.get("dosage_text"))]
    items = {
        str(product["id"]): {"basis_indication": "", **rule_based_versions(product.get("dosage_text"))}
        for product in products
    }
    level, reason = recommend_level_by_rule([p.get("dosage_text", "") for p in usable])
    return {
        "recommended_level": level,
        "reason": reason,
        "source": "원문 발췌" if usable else "요약할 원문 없음",
        "items": items,
        "warnings": [],
    }


def resolve_dosage_choice(choice, summaries):
    """화면에서 고른 값을 실제 버전 이름으로 바꿉니다. '자동(추천)'이면 추천 버전을 씁니다."""
    if choice == "원문" or not summaries:
        return "원문"
    if choice in DOSAGE_LEVELS:
        return choice
    return summaries.get("recommended_level", "일반")


def dosage_cell(item, level, original):
    """한 제품의 용법·용량 칸 내용을 만듭니다."""
    if level == "원문" or not item or is_status(original):
        return original
    text = item.get(DOSAGE_LEVEL_KEYS[level], "") or original
    basis = item.get("basis_indication", "")
    return f"({basis} 기준)\n{text}" if basis else text


# ── 사람이 채우는 빈 행 ─────────────────────────────────────

def manual_values(group_label=""):
    """한 제품 행에 넣을 빈 칸들. 신청의약품이면 원내코드 칸에 [신청의약품] 표시만 넣습니다."""
    values = {column: "" for column in MANUAL_COLUMNS}
    if group_label == "신청의약품":
        values["원내코드"] = APPLICANT_MARK
    return values


# ── Excel ─────────────────────────────────────────────────

def _styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="B7C3D0")
    return {
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "header_fill": PatternFill("solid", fgColor="DCE6F1"),
        "applicant_fill": PatternFill("solid", fgColor="FFF2CC"),
        "manual_fill": PatternFill("solid", fgColor="F2F2F2"),
        "wrap": Alignment(wrap_text=True, vertical="top"),
        "bold": Font(bold=True),
        "title": Font(bold=True, size=14),
        "status": Font(italic=True, color="808080"),
        "note": Font(size=9, color="595959"),
    }


def _write_sheet(sheet, df, title, note_lines, index_label, highlight_applicant):
    s = _styles()
    sheet["A1"] = title
    sheet["A1"].font = s["title"]
    start_row = 3
    headers = [index_label] + [str(column) for column in df.columns]
    for col_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=start_row, column=col_index, value=header)
        cell.font = s["bold"]
        cell.fill = s["applicant_fill"] if highlight_applicant and "신청의약품" in header else s["header_fill"]
        cell.alignment = s["wrap"]
        cell.border = s["border"]

    for row_offset, (label, row) in enumerate(df.iterrows(), start=1):
        excel_row = start_row + row_offset
        is_manual = str(label) in MANUAL_COLUMNS
        label_cell = sheet.cell(row=excel_row, column=1, value=str(label))
        label_cell.font = s["bold"]
        label_cell.alignment = s["wrap"]
        label_cell.border = s["border"]
        label_cell.fill = s["manual_fill"] if is_manual else s["header_fill"]
        for col_index, value in enumerate(row.tolist(), start=2):
            text = "" if value is None or (isinstance(value, float) and value != value) else str(value)
            cell = sheet.cell(row=excel_row, column=col_index, value=text)
            cell.alignment = s["wrap"]
            cell.border = s["border"]
            if is_manual:
                cell.fill = s["manual_fill"]
            if text in STATUS_VALUES:
                cell.font = s["status"]

    sheet.column_dimensions["A"].width = 22
    for col_index in range(2, len(headers) + 1):
        sheet.column_dimensions[sheet.cell(row=start_row, column=col_index).column_letter].width = 38
    sheet.freeze_panes = sheet.cell(row=start_row + 1, column=2)

    note_row = start_row + len(df) + 2
    for offset, line in enumerate(note_lines):
        sheet.cell(row=note_row + offset, column=1, value=line).font = s["note"]


# ── 비교표 셀 병합 (비교표 예시 모음.pptx 양식) ─────────────────
# - 같은 제품의 함량별 열은 맨 위 제품명 칸을 병합합니다.
# - 각 행에서 이웃한 칸의 내용이 같으면 병합하고, 다르면 나눠 둡니다.
#   같은 제품 안에서는 항상, 다른 제품끼리는 제품 경계가 딱 맞을 때만 병합합니다.
# - 함량마다 값이 따로 있어야 하는 행과 사람이 채우는 행은 병합하지 않습니다.
TABLE_HEADER_ROW = 3
NO_MERGE_ROWS = {
    "함량", "허가제품명", "영문제품명", "약가", "심평원제품코드", "약가적용시작일", "판매예정종료일", "원내코드",
} | set(MANUAL_TAIL_COLUMNS)
_STRENGTH_PATTERN = re.compile(
    r"\d+(?:\.\d+)?\s*(?:마이크로그램|밀리그램|그램|밀리리터|국제단위|단위|mcg|μg|mg|g|mL|ml|IU|%)"
    r"(?:\s*/\s*\d*(?:\.\d+)?\s*(?:밀리리터|mL|ml))?"
)


def base_product_name(name):
    """제품명에서 함량을 빼서 같은 제품인지 비교할 이름을 만듭니다. 예: 나르코설하정100마이크로그램 → 나르코설하정"""
    return _clean(_STRENGTH_PATTERN.sub("", str(name or "")))


def _table_rows(sheet):
    """비교표 표 부분의 (행 번호, 항목 이름) 목록. 표 아래 빈 줄부터는 설명 문구입니다."""
    rows, row = [], TABLE_HEADER_ROW + 1
    while sheet.cell(row=row, column=1).value not in (None, ""):
        rows.append((row, str(sheet.cell(row=row, column=1).value)))
        row += 1
    return rows


def unmerge_table(sheet):
    """표 안의 병합을 모두 풀고, 풀린 칸에 원래 값을 다시 채웁니다. (다시 병합하기 전에 씁니다)"""
    last_row = TABLE_HEADER_ROW + len(_table_rows(sheet))
    for merged in list(sheet.merged_cells.ranges):
        if merged.min_row < TABLE_HEADER_ROW or merged.min_row > last_row:
            continue
        value = sheet.cell(row=merged.min_row, column=merged.min_col).value
        sheet.unmerge_cells(str(merged))
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                sheet.cell(row=row, column=col).value = value


def product_groups(sheet):
    """열마다 어느 제품에 속하는지 정합니다. 같은 구분 + 같은 제품명(함량 제외)이면 한 제품입니다."""
    labels = {label: row for row, label in _table_rows(sheet)}
    name_row, group_row = labels.get("허가제품명"), labels.get("구분")
    keys = []
    for col in range(2, sheet.max_column + 1):
        name = sheet.cell(row=name_row, column=col).value if name_row else sheet.cell(row=TABLE_HEADER_ROW, column=col).value
        group = sheet.cell(row=group_row, column=col).value if group_row else ""
        keys.append((_clean(group), base_product_name(name)))
    return keys


def _merge_run(sheet, row, start_col, end_col):
    from openpyxl.styles import Alignment

    if end_col > start_col:
        sheet.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
        sheet.cell(row=row, column=start_col).alignment = Alignment(wrap_text=True, vertical="top", horizontal="center")


def apply_comparison_merges(sheet, cross_product=True):
    """비교표 시트에 병합 규칙을 적용합니다. 여러 번 실행해도 결과가 같습니다."""
    from openpyxl.styles import Font, PatternFill

    unmerge_table(sheet)
    keys = product_groups(sheet)
    if not keys:
        return
    first_col = 2

    # 맨 위 제품명 칸: 같은 제품의 함량별 열을 하나로 병합
    col = first_col
    while col < first_col + len(keys):
        end = col
        while end + 1 < first_col + len(keys) and keys[end + 1 - first_col] == keys[col - first_col]:
            end += 1
        group, name = keys[col - first_col]
        header = sheet.cell(row=TABLE_HEADER_ROW, column=col)
        header.value = f"{group} | {name}" if group else name
        header.font = Font(bold=True)
        if group == "신청의약품":
            header.fill = PatternFill("solid", fgColor="FFF2CC")
        _merge_run(sheet, TABLE_HEADER_ROW, col, end)
        col = end + 1

    # 각 항목 행: 이웃한 칸의 내용이 같으면 병합
    starts = {first_col + i for i, key in enumerate(keys) if i == 0 or keys[i - 1] != key}
    ends = {first_col + i for i, key in enumerate(keys) if i == len(keys) - 1 or keys[i + 1] != key}
    for row, label in _table_rows(sheet):
        if label in NO_MERGE_ROWS:
            continue
        values = [_clean(sheet.cell(row=row, column=c).value) for c in range(first_col, first_col + len(keys))]
        i = 0
        while i < len(values):
            j = i
            while j + 1 < len(values) and values[j + 1] == values[i]:
                j += 1
            if values[i]:
                run_start, run_end = first_col + i, first_col + j
                if cross_product and run_start in starts and run_end in ends:
                    _merge_run(sheet, row, run_start, run_end)
                else:
                    # 제품 경계가 맞지 않으면 같은 제품 안에서만 병합
                    sub_start = run_start
                    for c in range(run_start, run_end + 1):
                        if c == run_end or keys[c + 1 - first_col] != keys[c - first_col]:
                            _merge_run(sheet, row, sub_start, c)
                            sub_start = c + 1
            i = j + 1


def _write_source_sheet(sheet, dosage_sources):
    """용법·용량 원문 시트. 화면에서 어떤 버전을 골랐든 원문을 그대로 남깁니다."""
    s = _styles()
    for col_index, header in enumerate(SOURCE_HEADERS, start=1):
        cell = sheet.cell(row=1, column=col_index, value=header)
        cell.font = s["bold"]
        cell.fill = s["header_fill"]
        cell.border = s["border"]
    for row_index, source in enumerate(dosage_sources, start=2):
        for col_index, key in enumerate(SOURCE_KEYS, start=1):
            cell = sheet.cell(row=row_index, column=col_index, value=str(source.get(key, "") or (COMPARISON_SHEET if key == "sheet" else "")))
            cell.alignment = s["wrap"]
            cell.border = s["border"]
    for letter, width in zip("ABCDEF", [36, 14, 28, 50, 80, 14]):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A2"


def to_excel_bytes(comparison_df=None, detail_df=None, title="의약품 비교표", note_lines=None, dosage_sources=None):
    """비교표(제품=열, 항목=행), 상세 결과, 용법·용량 원문을 시트로 나눠 xlsx 파일 내용을 돌려줍니다.

    dosage_sources: [{"column": 비교표 열 이름, "id", "name", "indications_text", "dosage_text"}]
    """
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    notes = list(note_lines or []) + [
        "회색 행(원내코드, code open, 처방건수, 비급여)은 약사가 직접 입력합니다.",
        f"'{STATUS_FAILED}'는 API 호출이 실패한 칸, '{STATUS_EMPTY}'은 조회는 됐지만 내용이 없는 칸입니다.",
    ]
    if comparison_df is not None and len(comparison_df.columns):
        comparison_sheet = workbook.create_sheet(COMPARISON_SHEET)
        _write_sheet(comparison_sheet, comparison_df, title, notes, "항목", True)
        apply_comparison_merges(comparison_sheet)
    if detail_df is not None and len(detail_df):
        detail_t = detail_df.set_index("허가제품명").T if "허가제품명" in detail_df.columns else detail_df
        _write_sheet(workbook.create_sheet("상세결과"), detail_t, f"{title} — 상세 결과", notes, "항목", False)
    if dosage_sources:
        _write_source_sheet(workbook.create_sheet(SOURCE_SHEET), dosage_sources)
    if not workbook.sheetnames:
        workbook.create_sheet("비교표")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def comparison_sheet_names(workbook):
    """비교표 시트 이름들: '비교표' 하나, 또는 신청번호별 '신청1', '신청2' …"""
    if COMPARISON_SHEET in workbook.sheetnames:
        return [COMPARISON_SHEET]
    return [name for name in workbook.sheetnames if re.match(r"^신청\d+", name)]


def to_excel_bytes_by_number(parts, title="의약품 비교표", note_lines=None):
    """신청번호마다 비교표 시트를 따로 만듭니다: 시트 '신청1'에 신청의약품1과 그 비교약, '신청2'에 신청의약품2와 그 비교약 …

    parts: [{"number": "신청1", "name": 대표 제품명, "comparison": 비교표 DataFrame, "detail": 상세 DataFrame, "sources": 용법용량 원문}]
    상세결과와 용법용량 원문은 한 시트에 모으고 신청번호를 함께 적습니다.
    """
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    notes = list(note_lines or []) + [
        "회색 행(원내코드, code open, 처방건수, 비급여)은 약사가 직접 입력합니다.",
        f"'{STATUS_FAILED}'는 API 호출이 실패한 칸, '{STATUS_EMPTY}'은 조회는 됐지만 내용이 없는 칸입니다.",
    ]
    details, sources = [], []
    for part in parts:
        sheet_name = re.sub(r"[\\/*?:\[\]]", "", part["number"])[:31] or "신청"
        comparison = part.get("comparison")
        if comparison is not None and len(comparison.columns):
            sheet = workbook.create_sheet(sheet_name)
            _write_sheet(sheet, comparison, f"{title} — {part['number']} {part.get('name', '')}".strip(), notes, "항목", True)
            apply_comparison_merges(sheet)
        detail = part.get("detail")
        if detail is not None and len(detail):
            details.append(detail.assign(신청번호=part["number"])[["신청번호"] + list(detail.columns)])
        sources += [dict(source, sheet=sheet_name) for source in part.get("sources") or []]
    if details:
        import pandas as pd
        detail_all = pd.concat(details, ignore_index=True)
        detail_t = detail_all.set_index("허가제품명").T if "허가제품명" in detail_all.columns else detail_all
        _write_sheet(workbook.create_sheet("상세결과"), detail_t, f"{title} — 상세 결과 (신청번호 행 참고)", notes, "항목", False)
    if sources:
        _write_source_sheet(workbook.create_sheet(SOURCE_SHEET), sources)
    if not workbook.sheetnames:
        workbook.create_sheet(COMPARISON_SHEET)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
